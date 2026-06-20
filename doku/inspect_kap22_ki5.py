"""Inspects the KI5 section and surrounding chapters in the thesis."""
from docx import Document
import re

doc = Document(r'doku\HKTracker_CAS_Abschlussarbeit_Entwurf_v01.docx')

# Find KI5 and surrounding context
in_target = False
context_lines = []
target_paras = []

for i, para in enumerate(doc.paragraphs):
    txt = para.text.strip()
    style = para.style.name

    # Start capturing around KI5
    if 'KI' in txt and '5' in txt and ('Dokumenten' in txt or 'Beurteil' in txt or 'Evaluierung' in txt):
        in_target = True

    if in_target:
        target_paras.append((i, style, txt[:120]))
        if len(target_paras) > 80:
            break

# Also find Datenschutz section
print("=== KI5 SECTION ===")
for idx, style, txt in target_paras:
    print(f"[{idx:4d}] {style:<30} | {txt}")

print("\n=== SEARCHING FOR DATENSCHUTZ / TRANSPARENCY ===")
for i, para in enumerate(doc.paragraphs):
    txt = para.text.strip()
    if any(kw in txt for kw in ['Datenschutz', 'Transparenz', 'Azure', 'Mouseover', 'Feedback Loop', 'Glickman', 'rechtlich']):
        print(f"[{i:4d}] {para.style.name:<30} | {txt[:120]}")

print("\n=== SEARCHING FOR REFERENCES SECTION ===")
in_refs = False
for i, para in enumerate(doc.paragraphs):
    txt = para.text.strip()
    if 'Literatur' in txt or 'Referenz' in txt or 'Quellenverzeichnis' in txt:
        in_refs = True
    if in_refs:
        print(f"[{i:4d}] {para.style.name:<30} | {txt[:120]}")
        if len(txt) == 0 and i > 0:
            pass
        if in_refs and i > 0:
            pass

print("Done.")
