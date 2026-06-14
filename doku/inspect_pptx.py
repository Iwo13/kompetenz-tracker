import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation
from pptx.util import Pt

prs = Presentation(r'C:\Projekte\IT-Lehre Handlungskompetenzen\doku\Konzept_Kompetenzen_Ausbildungsplätze.pptx')
print(f'Slides: {len(prs.slides)}')
for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    texts.append(t)
    if texts:
        print(f'\n--- Slide {i+1} ---')
        for t in texts[:15]:
            print(f'  {t[:120]}')
