"""Verify the KI5 enrichment: check for new content and references."""
import sys
from docx import Document

doc = Document(r'doku\HKTracker_CAS_Abschlussarbeit_Entwurf_v01.docx')

checks = {
    'Datenschutzentscheid': False,
    'Mouseover-Tooltip': False,
    'Human-AI Feedback Loops': False,
    'R51': False,
    'R52': False,
    'R53': False,
    'Glickman': False,
    'EU AI Act': False,
    'FHNW Datenschutzrichtlinie': False,
}

for para in doc.paragraphs:
    t = para.text
    for key in checks:
        if key in t:
            checks[key] = True

print("=== Verification ===")
all_ok = True
for key, found in checks.items():
    status = "OK" if found else "MISSING"
    print(f"  [{status}] {key}")
    if not found:
        all_ok = False

print()
print("=== KI5 Sub-bullets (around Human-Kontrolle) ===")
in_ki5 = False
for i, para in enumerate(doc.paragraphs):
    if 'KI 5:' in para.text:
        in_ki5 = True
    if in_ki5 and ('LiteLLM' in para.text or 'Nicht eingesetzte' in para.text):
        break
    if in_ki5:
        print(f"  [{i:4d}] {para.text[:120]}")

print()
print("=== References R49-R53 ===")
for i, para in enumerate(doc.paragraphs):
    if any(f'[R{n}]' in para.text for n in [49, 50, 51, 52, 53]):
        print(f"  [{i:4d}] {para.text[:120]}")
