import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = r"doku\HKTracker_CAS_Abschlussarbeit_Entwurf_v01.docx"
doc = Document(DOC)
paras = doc.paragraphs

list_para_style = doc.styles['List Paragraph']

# ── 1. BB und PPB Stil korrigieren ────────────────────────────────
for p in paras:
    if p.text.startswith("BB:\t") or p.text.startswith("PPB:\t"):
        p.style = list_para_style
        print(f"  [FIX]  Stil korrigiert: {p.text[:50]}")

# ── 2. L hinzufügen falls noch nicht vorhanden ────────────────────
# Prüfe ob L bereits korrekt vorhanden (als eigener Eintrag)
l_exists = any(
    p.text.startswith("L:\t") and 'Lernende' in p.text
    for p in paras
)

if l_exists:
    print("  [SKIP] L bereits vorhanden")
else:
    # Einfügen nach KI / AI (vor LLM)
    ref = next((p for p in paras if p.text.startswith("KI / AI:\t")), None)
    if not ref:
        print("  [FEHLER] KI/AI-Eintrag nicht gefunden")
    else:
        # Neuen Paragraph mit korrektem Stil erstellen
        new_p = doc.add_paragraph(
            "L:\tLernende/r (Rollenkürzel im Anforderungskatalog; Persona Heinz/Viktor)",
            style='List Paragraph'
        )
        # Verschieben: aus Ende des Dokuments an die richtige Stelle
        ref._p.addnext(new_p._p)
        print(f"  [OK]   L eingefügt nach: {ref.text[:40]}")

doc.save(DOC)
print(f"\nGespeichert: {DOC}")
