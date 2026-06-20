"""
Generates a Gantt chart PNG for the HK-Tracker development timeline
and inserts it into the thesis DOCX at the [TODO: Zeitplan] paragraph.

Timeline:
  Phase 1 (abgeschlossen): März–Juni 2026
  Phase 2 (KI-Integration, geplant): Juli–August 2026
  Phase 3 (Pilotbetrieb, geplant): Juli–August 2026
  Thesis-Abgabe: 16. August 2026
"""
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import matplotlib.dates as mdates
from matplotlib.lines import Line2D
from datetime import datetime, timedelta
from xml.sax.saxutils import escape as xml_escape
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

DOCX_PATH = r'doku\HKTracker_CAS_Abschlussarbeit_Entwurf_v01.docx'
IMG_PATH  = r'doku\gantt_zeitplan.png'
TODAY     = datetime(2026, 6, 20)
DEADLINE  = datetime(2026, 8, 16)

# FHNW Colors
NAVY   = '#002B5C'
YELLOW = '#FFD700'
GREEN  = '#16a34a'    # done tasks
BLUE   = '#3b82f6'    # planned tasks
LGRAY  = '#f8fafc'
MGRAY  = '#e2e8f0'
DGRAY  = '#94a3b8'

def dt(month, day):
    return datetime(2026, month, day)

# -------------------------------------------------------------------------
# Task definitions
# Each: (label, start, end, done, section_idx)
# -------------------------------------------------------------------------
SECTIONS = [
    ('Phase 1', 'Entwicklung (abgeschlossen)', '#dcfce7', '#166534'),
    ('Phase 2', 'KI-Integration (geplant)',    '#dbeafe', '#1d4ed8'),
    ('Phase 3', 'Pilotbetrieb (geplant)',       '#fff7ed', '#c2410c'),
]

TASKS = [
    # Phase 1 — DONE
    ('Stack-Migration & UI-Grundgerüst',        dt(3, 1),  dt(4, 30), True,  0),
    ('Backend-Migration ASP.NET Core 8',         dt(4, 1),  dt(5, 31), True,  0),
    ('Kernfunktionen + Rotationsplanung Gantt',  dt(5, 1),  dt(5, 31), True,  0),
    ('KI5 Dokumentenbeurteilung (impl.)',        dt(5,15),  dt(6, 20), True,  0),
    # Phase 2 — PLANNED
    ('Azure OpenAI Integration (Key + Test)',    dt(7, 1),  dt(7, 31), False, 1),
    ('MSAL-Authentifizierung (IT-Freigabe)',     dt(7, 1),  dt(7, 31), False, 1),
    ('Deployment Azure App Service',             dt(7,15),  dt(8, 15), False, 1),
    # Phase 3 — PLANNED
    ('Early Adopter Onboarding',                 dt(7,14),  dt(7, 21), False, 2),
    ('Pilotbetrieb (4 Wochen)',                  dt(7,14),  dt(8, 10), False, 2),
    ('Datenerhebung & Auswertung',               dt(8, 1),  dt(8, 16), False, 2),
]

# -------------------------------------------------------------------------
# Build row list (section headers + tasks, top→bottom)
# -------------------------------------------------------------------------
rows = []
prev_sec = -1
for task in TASKS:
    sec = task[4]
    if sec != prev_sec:
        rows.append(('header', SECTIONS[sec], sec))
        prev_sec = sec
    rows.append(('task', task))

n = len(rows)
# y positions: row 0 = top → y = n; row n-1 = bottom → y = 1
def y_pos(i):
    return n - i

# -------------------------------------------------------------------------
# Draw
# -------------------------------------------------------------------------
fig_w, fig_h = 13, 0.55 * n + 1.8
fig, ax = plt.subplots(figsize=(fig_w, fig_h))
fig.patch.set_facecolor('white')
ax.set_facecolor(LGRAY)

XMIN = datetime(2026, 2, 25)
XMAX = datetime(2026, 9, 5)
ax.set_xlim(XMIN, XMAX)
ax.set_ylim(0.2, n + 0.8)

bar_h       = 0.52
header_h    = 0.62
task_color  = {True: GREEN, False: BLUE}

ytick_pos, ytick_lbl = [], []

for i, row in enumerate(rows):
    yp = y_pos(i)
    if row[0] == 'header':
        _, sec_info, sec_idx = row
        short, long, bg, fg = sec_info
        # Full-width header band
        ax.barh(yp, XMAX - XMIN, left=XMIN, height=header_h,
                color=NAVY, alpha=1.0, zorder=3)
        ax.text(XMIN + timedelta(days=2), yp,
                f'  {short} – {long}',
                va='center', ha='left', fontsize=8.5, fontweight='bold',
                color='white', zorder=4)
    else:
        _, task = row
        label, start, end, done, sec_idx = task
        color = task_color[done]
        bg = SECTIONS[sec_idx][2]

        # Subtle row background
        ax.barh(yp, XMAX - XMIN, left=XMIN, height=0.75,
                color=bg, alpha=0.35, zorder=1)
        # Task bar
        duration = end - start
        ax.barh(yp, duration, left=start, height=bar_h,
                color=color, alpha=0.90, zorder=3,
                edgecolor='white', linewidth=0.8)
        # Checkmark for done
        if done:
            mid = start + duration / 2
            ax.text(mid, yp, '✓', va='center', ha='center',
                    fontsize=8, color='white', fontweight='bold', zorder=4)
        # Start/end date labels (subtle)
        ax.text(start - timedelta(days=1), yp, start.strftime('%d.%m'),
                va='center', ha='right', fontsize=6.5, color=DGRAY, zorder=4)
        ax.text(end + timedelta(days=1), yp, end.strftime('%d.%m'),
                va='center', ha='left', fontsize=6.5, color=DGRAY, zorder=4)

        ytick_pos.append(yp)
        ytick_lbl.append(label)

# X-axis
ax.xaxis.set_major_locator(mdates.MonthLocator())
ax.xaxis.set_major_formatter(mdates.DateFormatter('%b %Y'))
ax.xaxis.set_minor_locator(mdates.WeekdayLocator(byweekday=0, interval=2))
plt.xticks(rotation=35, ha='right', fontsize=8.5)
ax.tick_params(axis='x', which='minor', length=3, color=MGRAY)

# Y-axis
ax.set_yticks(ytick_pos)
ax.set_yticklabels(ytick_lbl, fontsize=8.5)
ax.tick_params(axis='y', length=0)

# Vertical grid (months)
ax.xaxis.grid(True, which='major', color=MGRAY, linewidth=0.8, zorder=0)
ax.xaxis.grid(True, which='minor', color=MGRAY, linewidth=0.3, zorder=0)
ax.set_axisbelow(True)

# Spines
for spine in ['top', 'right', 'left']:
    ax.spines[spine].set_visible(False)
ax.spines['bottom'].set_color(MGRAY)

# TODAY line
ax.axvline(TODAY, color='#dc2626', linewidth=2, linestyle='--', zorder=5)
ax.text(TODAY + timedelta(days=1.5), n + 0.55, 'Heute\n20.06.',
        color='#dc2626', fontsize=7.5, fontweight='bold', va='top', zorder=6)

# DEADLINE line
ax.axvline(DEADLINE, color=NAVY, linewidth=2.5, linestyle='-', zorder=5)
ax.text(DEADLINE + timedelta(days=1.5), n + 0.55,
        'Thesis-Abgabe\n16.08.2026',
        color=NAVY, fontsize=7.5, fontweight='bold', va='top', zorder=6)

# Legend
legend_elements = [
    mpatches.Patch(facecolor=GREEN,   edgecolor='white', label='Abgeschlossen ✓'),
    mpatches.Patch(facecolor=BLUE,    edgecolor='white', label='Geplant'),
    Line2D([0], [0], color='#dc2626', linewidth=2, linestyle='--', label='Heute (20.06.2026)'),
    Line2D([0], [0], color=NAVY,      linewidth=2.5,     label='Thesis-Abgabe (16.08.2026)'),
]
ax.legend(handles=legend_elements, loc='lower right', fontsize=8,
          framealpha=0.95, frameon=True, edgecolor=MGRAY)

plt.title('HK-Tracker – Entwicklungs- und Einführungsplan',
          fontsize=11, fontweight='bold', color=NAVY, pad=12)
plt.tight_layout(pad=1.5)
plt.savefig(IMG_PATH, dpi=150, bbox_inches='tight', facecolor='white')
plt.close()
print(f"PNG saved: {IMG_PATH}")


# -------------------------------------------------------------------------
# Insert into DOCX
# -------------------------------------------------------------------------
def make_para(text, style_id):
    return etree.fromstring(
        f'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr><w:pStyle w:val="{style_id}"/>'
        f'<w:jc w:val="center"/>'
        f'</w:pPr>'
        f'<w:r><w:t xml:space="preserve">{xml_escape(text)}</w:t></w:r>'
        f'</w:p>'
    )

doc = Document(DOCX_PATH)

todo_para = None
for para in doc.paragraphs:
    if '[TODO: Zeitplan als Gantt' in para.text:
        todo_para = para
        break

if todo_para is None:
    print("WARNING: [TODO: Zeitplan] not found — appending image at end of document")
    img_para = doc.add_paragraph()
    img_para.alignment = 1
    run = img_para.add_run()
    run.add_picture(IMG_PATH, width=Inches(5.8))
else:
    print(f"Found TODO para: {todo_para.text[:80]}")

    # 1. Add image paragraph at end of doc (python-docx easiest insertion point)
    img_para = doc.add_paragraph()
    img_para.alignment = 1  # CENTER
    run = img_para.add_run()
    run.add_picture(IMG_PATH, width=Inches(5.8))

    # 2. Move image paragraph to just BEFORE the TODO para
    todo_para._p.addprevious(img_para._p)

    # 3. Replace TODO text with caption
    for t in todo_para._p.iter(qn('w:t')):
        t.text = ''
    # Clear runs
    for r in todo_para._p.findall(qn('w:r')):
        todo_para._p.remove(r)
    # Add caption text
    caption_text = ('Abbildung: HK-Tracker Entwicklungs- und Einführungsplan '
                    '(März–August 2026) — Grün = abgeschlossen; Blau = geplant; '
                    'Abgabe Abschlussarbeit 16.08.2026')
    r_elem = etree.fromstring(
        f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:rPr><w:i/><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>'
        f'<w:t xml:space="preserve">{xml_escape(caption_text)}</w:t>'
        f'</w:r>'
    )
    # Center-align the caption paragraph
    pPr = todo_para._p.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        todo_para._p.insert(0, pPr)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    pPr.append(jc)
    todo_para._p.append(r_elem)

    print("OK: Image inserted and TODO replaced with caption")

doc.save(DOCX_PATH)
print(f"Saved: {DOCX_PATH}")

print()
print("=== MERMAID CODE (for GitHub / web rendering) ===")
print("""
```mermaid
gantt
    title HK-Tracker – Entwicklungs- und Einführungsplan
    dateFormat  YYYY-MM-DD
    axisFormat  %b %Y

    section Phase 1 – Entwicklung (abgeschlossen)
    Stack-Migration & UI-Grundgerüst        :done, p1a, 2026-03-01, 2026-04-30
    Backend-Migration ASP.NET Core 8        :done, p1b, 2026-04-01, 2026-05-31
    Kernfunktionen + Rotationsplanung Gantt :done, p1c, 2026-05-01, 2026-05-31
    KI5 Dokumentenbeurteilung (impl.)       :done, p1d, 2026-05-15, 2026-06-20

    section Phase 2 – KI-Integration (geplant)
    Azure OpenAI Integration                :active, p2a, 2026-07-01, 2026-07-31
    MSAL-Authentifizierung (IT-Freigabe)    :p2b, 2026-07-01, 2026-07-31
    Deployment Azure App Service            :p2c, 2026-07-15, 2026-08-15

    section Phase 3 – Pilotbetrieb (geplant)
    Early Adopter Onboarding                :p3a, 2026-07-14, 2026-07-21
    Pilotbetrieb (4 Wochen)                :p3b, 2026-07-14, 2026-08-10
    Datenerhebung & Auswertung             :p3c, 2026-08-01, 2026-08-16

    section Meilensteine
    Thesis-Abgabe                           :milestone, m1, 2026-08-16, 0d
```
""")
